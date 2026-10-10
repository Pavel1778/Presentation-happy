"""Audit the delivered PPTX files against the deck requirements.

Checks structure, notes, absence of audio/video, transitions, animations,
Morph wiring, and QR decodability. Exits non-zero if any check fails.
"""
import re
import sys
import zipfile
from pathlib import Path

import cv2
import numpy as np
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[4]


class Audit:
    def __init__(self) -> None:
        self.ok = True

    def check(self, name: str, cond: bool, extra: object = '') -> None:
        self.ok = self.ok and cond
        print(('  OK   ' if cond else ' FAIL  '), name, extra)


def audit(path: Path) -> Audit:
    a = Audit()
    print('=' * 60)
    print(path.name)
    pres = Presentation(str(path))
    z = zipfile.ZipFile(path)
    names = z.namelist()

    a.check('17 slides', len(pres.slides) == 17, len(pres.slides))
    a.check('16:9 1280x720',
            (pres.slide_width, pres.slide_height) == (12192000, 6858000),
            f'{pres.slide_width}x{pres.slide_height}')
    notes = sum(1 for s in pres.slides
                if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    a.check('notes on 17 slides', notes == 17, notes)
    a.check('no audio parts', not any(n.endswith(('.mp3', '.wav', '.m4a')) for n in names))
    a.check('no mp4 parts', not any(n.endswith('.mp4') for n in names))
    xml = ''.join(z.read(n).decode('utf-8', 'ignore')
                  for n in names if n.endswith(('.xml', '.rels')))
    a.check('no audio/video wiring',
            all(t not in xml for t in
                ('audioFile', 'ppaction://media', 'p14:media', 'videoFile')))
    a.check('animated GIF present', any(n.endswith('.gif') for n in names))

    fade = morph = 0
    anim = morph_pairs = 0
    for i in range(1, 18):
        s = z.read(f'ppt/slides/slide{i}.xml').decode()
        if 'p159:morph' in s:
            morph += 1
        elif '<p:fade' in s:
            fade += 1
        if '<p:timing>' in s:
            anim += 1
        if '!!title-flow' in s:
            morph_pairs += 1
    a.check('transitions 13 fade / 3 morph / 1 none',
            fade == 13 and morph == 3, f'fade={fade} morph={morph}')
    a.check('object animation on 16 slides', anim == 16, anim)
    a.check('morph pairs on 6 slides', morph_pairs == 6, morph_pairs)
    a.check('no speaker-note audio', xml.count('audioFile') == 0)

    rels = z.read('ppt/slides/_rels/slide17.xml.rels').decode()
    m = re.search(r'media/([^"]+\.png)', rels)
    a.check('slide-17 QR is a PNG part', m is not None)
    if m:
        png = z.read('ppt/media/' + m.group(1))
        dec = cv2.QRCodeDetector().detectAndDecode(
            cv2.imdecode(np.frombuffer(png, np.uint8), cv2.IMREAD_GRAYSCALE))[0]
        a.check('QR decodes',
                dec.startswith('https://github.com/Pavel1778/Presentation-happy'),
                dec[:60])

    print('RESULT', 'PASS' if a.ok else 'FAIL')
    return a


if __name__ == '__main__':
    targets = sys.argv[1:] or [
        str(ROOT / 'Konceptsiya_schastya.pptx'),
        str(ROOT / 'Konceptsiya_schastya_bez_video.pptx'),
    ]
    all_ok = all(audit(Path(t)).ok for t in targets)
    sys.exit(0 if all_ok else 1)
